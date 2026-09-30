"""
Rule-based budget and run-rate forecast: the template for any company.

Each P&L account is budgeted by one rule (company.toml [budget]):

    growth       prior-year same month x (1 + rate)          keeps last year's seasonality
    prior_year   prior-year same month
    pct_revenue  rate x total budget revenue of the month    e.g. materials, card fees
    fixed        a monthly amount ("monthly")                e.g. rent, software
    manual       twelve monthly amounts ("months")
    tax          rate x profit before tax of the month

Defaults per section apply to accounts without their own rule:

    [budget.defaults]
    revenue = {method = "growth", rate = 0.05}
    opex = {method = "growth", rate = 0.03}
    [budget.accounts]
    "310" = {method = "pct_revenue", rate = 0.12, note = "materials follow sales"}

The forecast for the open months scales the remaining budget by the year-to-date
run-rate (actual / budget) of each section, as set in forecast.toml:

    [latest_estimate]
    revenue = "ytd"          # or a factor, e.g. 0.97
    cost_of_sales = "ytd"
    opex = "budget"          # = 1.0
    adjustments = {"477" = -400.0}      # monthly change for named accounts
    [scenarios.downside]
    revenue_mult = 0.92
"""
from . import company, ledger

METHODS = ("growth", "prior_year", "pct_revenue", "fixed", "manual", "tax")
DEFAULTS = {"revenue": {"method": "growth", "rate": 0.05}, "cost_of_sales": {"method": "growth", "rate": 0.03},
            "opex": {"method": "growth", "rate": 0.03}, "depreciation": {"method": "prior_year"},
            "interest": {"method": "prior_year"}, "tax": {"method": "tax", "rate": 0.25}}


class RulesError(ValueError):
    pass


def rules(co):
    cfg = co.cfg.get("budget", {})
    defaults = {**DEFAULTS, **cfg.get("defaults", {})}
    out = {}
    for a in co.chart.accounts:
        r = dict(defaults[a.section])
        r.update(cfg.get("accounts", {}).get(a.code, {}))
        if r.get("method") not in METHODS:
            raise RulesError("account %s: unknown budget method %r (use %s)" % (a.code, r.get("method"), ", ".join(METHODS)))
        if r["method"] == "pct_revenue" and a.section == "revenue":
            raise RulesError("account %s: a revenue account cannot be a percentage of revenue" % a.code)
        if r["method"] == "manual" and len(r.get("months", [])) != 12:
            raise RulesError("account %s: manual budgets need 12 monthly amounts" % a.code)
        out[a.code] = r
    return out


def prior_year(co, root=None):
    months = co.fy_months(co.budget_year - 1)
    have = set(ledger.available_months(co, root))
    missing = [m for m in months if m not in have]
    return months, missing


def budget(co, months, root=None):
    """{month: {code: amount}} for the budget year."""
    rs = rules(co)
    py_months, missing = prior_year(co, root)
    needs_py = any(r["method"] in ("growth", "prior_year") for r in rs.values())
    if needs_py and missing:
        raise RulesError("the budget grows from last year's actuals: load %s first (missing %s)"
                         % (co.fy_label(co.budget_year - 1), ", ".join(missing)))
    py = {m: ledger.read_month(co, m, root)[0] for m in py_months if m not in missing}
    chart = co.chart
    out = {}
    for i, m in enumerate(months):
        p = {}
        pm = py_months[i]

        def base(code, r):
            meth = r["method"]
            if meth == "growth":
                return py[pm][code] * (1 + r.get("rate", 0.0))
            if meth == "prior_year":
                return py[pm][code]
            if meth == "fixed":
                return r.get("monthly", r.get("annual", 0.0) / 12)
            if meth == "manual":
                return r["months"][i]
            return None

        for code in chart.codes("revenue"):
            p[code] = base(code, rs[code])
        rev = sum(p.values())
        for code in chart.codes("cost_of_sales", "opex", "depreciation", "interest"):
            r = rs[code]
            p[code] = r["rate"] * rev if r["method"] == "pct_revenue" else base(code, r)
            if p[code] is None:
                raise RulesError("account %s: method %s is only for tax accounts" % (code, r["method"]))
        pbt = chart.totals(p)["pbt"]
        for code in chart.codes("tax"):
            r = rs[code]
            p[code] = r.get("rate", 0.25) * pbt if r["method"] == "tax" else base(code, r)
        out[m] = {c: round(v, 2) for c, v in p.items()}
    return out


# ---------------------------------------------------------------- Excel
def budget_sheets(wb, co, months, root=None):
    """Prior year, Assumptions (rules) and Budget sheets with live formulas."""
    from . import workbook as W
    from openpyxl.styles import Alignment
    chart = co.chart
    rs = rules(co)
    py_months, missing = prior_year(co, root)
    py = {m: ledger.read_month(co, m, root)[0] for m in py_months if m not in missing}

    ws = wb.create_sheet("Prior year")
    W._title(ws, "Prior year actuals: %s" % co.fy_label(co.budget_year - 1),
             "%s. GBP, ex VAT. The base for the growth rules of the budget." % co.name)
    W._index_row(ws, 3)
    W._month_header(ws, 4, py_months, co)
    ws.cell(row=4, column=W.FIRST_COL + 12, value=co.fy_label(co.budget_year - 1))
    W.write_pnl(ws, co, py_months, lambda code, i, c: py[py_months[i]][code] if py_months[i] in py else None,
                show=lambda i: py_months[i] in py)
    ws.freeze_panes = "D5"

    wa = wb.create_sheet("Assumptions")
    W._title(wa, "Budget rules: %s" % co.fy_label(), "%s. One rule per account; blue = input." % co.name)
    heads = ["Account", "Code", "Section", "Method", "Rate", "Monthly", "Note"]
    widths = [34, 8, 14, 12, 9, 11, 60]
    for j, (h, w) in enumerate(zip(heads, widths), start=1):
        c = wa.cell(row=4, column=j, value=h)
        c.fill, c.font = W.FILL_HEAD, W.F_HEAD
        wa.column_dimensions[W.get_column_letter(j)].width = w
    arow = {}
    row = 5
    for a in chart.accounts:
        r = rs[a.code]
        arow[a.code] = row
        wa.cell(row=row, column=1, value=a.name).font = W.F_BASE
        wa.cell(row=row, column=2, value=a.code).font = W.F_NOTE
        wa.cell(row=row, column=3, value=a.section.replace("_", " ")).font = W.F_NOTE
        wa.cell(row=row, column=4, value=r["method"]).font = W.F_INPUT
        if r["method"] in ("growth", "pct_revenue", "tax"):
            c = wa.cell(row=row, column=5, value=r.get("rate", 0.25 if r["method"] == "tax" else 0.0))
            c.font, c.number_format = W.F_INPUT, W.PCT
        if r["method"] == "fixed":
            c = wa.cell(row=row, column=6, value=r.get("monthly", r.get("annual", 0.0) / 12))
            c.font, c.number_format = W.F_INPUT, W.GBP
        wa.cell(row=row, column=7, value=r.get("note", "")).font = W.F_NOTE
        row += 1
    wa.cell(row=row + 1, column=1, value="Methods: growth = prior-year month x (1 + rate); prior_year; pct_revenue = "
                                         "rate x budget revenue; fixed = monthly amount; manual; tax = rate x profit "
                                         "before tax.").font = W.F_NOTE
    wa.freeze_panes = "A5"

    r = W.pnl_rows(chart, 5)
    base_row = r["pbt"] if chart.has_interest else r["ebit"]

    def value(code, i, c):
        rule, ar = rs[code], arow[code]
        meth = rule["method"]
        if meth == "growth":
            return "='Prior year'!%s%d*(1+Assumptions!$E$%d)" % (c, r["L" + code], ar)
        if meth == "prior_year":
            return "='Prior year'!%s%d" % (c, r["L" + code])
        if meth == "pct_revenue":
            return "=Assumptions!$E$%d*%s%d" % (ar, c, r["revenue"])
        if meth == "fixed":
            return "=Assumptions!$F$%d" % ar
        if meth == "manual":
            return rule["months"][i]
        return "=Assumptions!$E$%d*%s%d" % (ar, c, base_row)

    b = W.budget_sheet(wb, co, months, value, "Calculated from the budget rules (Assumptions) and the prior year.")
    wsb = wb["Budget"]
    for code, rule in rs.items():
        if rule["method"] == "manual":
            for i in range(12):
                wsb.cell(row=r["L" + code], column=W.FIRST_COL + i).font = W.F_INPUT
    return {"assumptions": arow, "prior_year": r, "budget": b}


# ---------------------------------------------------------------- forecast
def _factor(setting, act_sum, bud_sum):
    if setting in (None, "budget"):
        return 1.0
    if setting == "ytd":
        return min(1.5, max(0.5, act_sum / bud_sum)) if bud_sum else 1.0
    return float(setting)


def reforecast(co, b, act, closed, open_, le, scenario=None):
    chart = co.chart
    est = dict(le.get("latest_estimate", {}))
    notes = est.pop("notes", {})
    adjustments = est.pop("adjustments", {})
    sc = (le.get("scenarios", {}).get(scenario) if scenario else None) or {}
    if scenario and not sc:
        from .forecast import ForecastError
        raise ForecastError("unknown scenario %s" % scenario)
    sec_sum = lambda src, months, sec: sum(src[m][c] for m in months for c in chart.codes(sec))
    act_p = {m: act[m]["pnl"] for m in closed}
    factors = {}
    for sec in ("revenue", "cost_of_sales", "opex"):
        f = _factor(est.get(sec, "ytd" if sec != "opex" else "budget"), sec_sum(act_p, closed, sec),
                    sec_sum(b["pnl"], closed, sec)) if closed else 1.0
        mult = sc.get("revenue_mult", 1.0) if sec == "revenue" else sc.get("cost_mult", 1.0)
        if sec == "cost_of_sales":
            mult *= sc.get("revenue_mult", 1.0)          # direct costs move with sales
        factors[sec] = f * mult
    tax_rate = {c: (co.cfg.get("budget", {}).get("accounts", {}).get(c, {}).get("rate")
                    or co.cfg.get("budget", {}).get("defaults", {}).get("tax", {}).get("rate", 0.25))
                for c in chart.codes("tax")}
    by_month = {m: act_p[m] for m in closed}
    for m in open_:
        p = {}
        for a in chart.accounts:
            v = b["pnl"][m][a.code]
            if a.section in factors:
                v *= factors[a.section]
            p[a.code] = v + adjustments.get(a.code, 0.0)
        pbt = chart.totals(p)["pbt"]
        for c in chart.codes("tax"):
            p[c] = tax_rate[c] * pbt
        by_month[m] = {c: round(v, 2) for c, v in p.items()}
    fy = co.fy_months()
    fy_pnl = {c: sum(by_month[m][c] for m in fy) for c in chart.by_code}
    bfy = {c: sum(b["pnl"][m][c] for m in fy) for c in chart.by_code}
    assumptions = {"revenue_factor": factors["revenue"], "opex_factor": factors["opex"]}
    if chart.has_cost_of_sales:
        assumptions["cost_of_sales_factor"] = factors["cost_of_sales"]
    for code, v in adjustments.items():
        assumptions["adjustment_%s" % code] = v
    return {
        "label": "%d+%d" % (len(closed), len(open_)), "scenario": scenario or "base",
        "months": fy, "closed": closed, "open": open_, "pnl": by_month,
        "source": {m: ("A" if m in act else "F") for m in fy},
        "totals": {m: chart.totals(by_month[m]) for m in fy},
        "fy": fy_pnl, "fy_totals": chart.totals(fy_pnl), "budget_fy": bfy, "budget_fy_totals": chart.totals(bfy),
        "drivers": [], "assumptions": assumptions,
        "assumption_notes": {"revenue_factor": notes.get("revenue", "Year-to-date revenue vs budget carried forward."),
                             "cost_of_sales_factor": notes.get("cost_of_sales", "Direct costs follow the year-to-date run-rate."),
                             "opex_factor": notes.get("opex", "Overheads as budgeted.")},
    }

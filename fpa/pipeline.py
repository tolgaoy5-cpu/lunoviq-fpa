"""
Monthly run: builds the management pack for a reporting month, recalculates it
in Excel, audits Excel against the Python engine and writes summary.json for
the web panel.

    output/<company>/<YYYY-MM>_<timestamp>/<Company>_FPA_<YYYY-MM>.xlsx
    output/<company>/<YYYY-MM>_<timestamp>/summary.json
"""
import datetime as dt
import json
import re
from pathlib import Path

import openpyxl

from . import company, workbook
from .forecast import load_le

OUTPUT = company.ROOT / "output"
STEPS = ["actuals", "workbook", "recalc", "audit"]


def audit(co, path, lay, root=None):
    """Excel values vs the Python engine; returns counts and any mismatches."""
    wb = openpyxl.load_workbook(path, data_only=True)
    miss, n = [], 0

    def check(where, got, want, tol):
        nonlocal n
        n += 1
        if got is None or abs(float(got) - want) > tol:
            miss.append({"where": where, "excel": got, "python": round(want, 2)})

    from .budget import build as budget_build
    b = budget_build(co, root=root)
    ws = wb["Budget"]
    for i, m in enumerate(lay["months"]):
        for code in co.chart.by_code:
            check("Budget %s %s" % (m, code), ws["%s%d" % (workbook.mcol(i), lay["budget"]["L" + code])].value,
                  b["pnl"][m][code], 0.1)
    ws, res = wb["BvA"], lay["variance"]
    for code in co.chart.by_code:
        r = lay["bva"][code]
        for col, per, f in (("D", "month", "actual"), ("E", "month", "budget"), ("I", "ytd", "actual"),
                            ("J", "ytd", "budget"), ("K", "ytd", "var")):
            check("BvA %s %s %s" % (code, per, f), ws["%s%d" % (col, r)].value, res["lines"][per][code][f], 0.1)
    ws, fc = wb["Forecast"], lay["scenario_results"]["base"]
    for code in co.chart.by_code:
        check("Forecast FY %s" % code, ws["%s%d" % (workbook.TOTAL, lay["forecast"]["pnl"]["L" + code])].value,
              fc["fy"][code], 0.1)
    ws, cf = wb["Cash13W"], lay["cash_result"]
    for i in range(13):
        check("Cash week %d" % (i + 1), ws.cell(row=lay["cash"]["closing"], column=workbook.FIRST_COL + i).value,
              cf["closing"][i], 0.5)
    errors = [(s.title, c.coordinate) for s in wb for row in s.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#") and c.value[1:].rstrip("!?/0").isupper()]
    overall = wb["Checks"].cell(row=lay["checks"]["overall"], column=4).value
    return {"values_checked": n, "mismatches": miss, "excel_errors": errors, "model_checks": overall}


def summary(co, lay, audit_result, files, root=None):
    res, fc, sc, cf = lay["variance"], lay["scenario_results"]["base"], lay["scenario_results"], lay["cash_result"]

    def lines(period):
        out = []
        for kind, key, label in workbook.pnl_layout(co.chart):
            if kind == "account":
                l = res["lines"][period][key]
                a = co.chart.by_code[key]
                out.append({"code": key, "label": label, "group": a.line, "section": a.section, **l})
        return out

    from .budget import build as budget_build
    bud = budget_build(co, root=root)
    return {
        "company": co.name, "slug": co.slug, "model": co.model, "demo": bool(co.cfg["company"].get("demo")),
        "month": res["month"], "fy": co.fy_label(),
        "forecast_label": fc["label"], "has_cost_of_sales": co.chart.has_cost_of_sales,
        "reporting": co.cfg.get("reporting", {}),
        "chart_lines": [{"line": l, "section": s} for l, s in co.chart.lines("revenue", "cost_of_sales", "opex")],
        "generated": dt.datetime.now().isoformat(timespec="seconds"),
        "totals": {p: {k: v for k, v in res["totals"][p].items()} for p in ("month", "ytd")},
        "lines": {p: lines(p) for p in ("month", "ytd")},
        "commentary": {p: [{"label": l, "text": t, "favourable": f} for l, t, f in res["commentary"][p]]
                       for p in ("month", "ytd")},
        "bridges": res["bridges"],
        "kpis": [{"label": l, "actual": a, "budget": b_, "kind": k} for l, a, b_, k in res["kpis"]],
        "ytd_kpis": [{"label": l, "actual": a, "budget": b_} for l, a, b_ in res["ytd_kpis"]],
        "monthly": [{"month": m, "source": fc["source"][m], "revenue": fc["totals"][m]["revenue"],
                     "gross_profit": fc["totals"][m]["gross_profit"],
                     "budget_gross_profit": bud["totals"][m]["gross_profit"],
                     "ebitda": fc["totals"][m]["ebitda"], "budget_revenue": bud["totals"][m]["revenue"],
                     "budget_ebitda": bud["totals"][m]["ebitda"]} for m in fc["months"]],
        "full_year": {"budget": fc["budget_fy_totals"],
                      **{n: r["fy_totals"] for n, r in sc.items()},
                      "closing_pum": {n: r.get("closing_pum") for n, r in sc.items()},
                      "budget_closing_pum": fc.get("budget_closing_pum")},
        "assumptions": fc["assumptions"],
        "assumption_notes": fc.get("assumption_notes") or load_le(co).get("latest_estimate", {}).get("notes", {}),
        "cash": {"weeks": [w0.isoformat() for w0, _ in cf["weeks"]], "receipts": cf["receipts"],
                 "payments": cf["payments"], "closing": cf["closing"], "opening": cf["opening"],
                 "minimum": cf["minimum"], "lowest": cf["lowest"], "lowest_week": cf["lowest_week"],
                 "lines": {"%s|%s" % k: v for k, v in cf["table"].items()},
                 "corporation_tax": {**cf["corporation_tax"], "due": cf["corporation_tax"]["due"].isoformat()},
                 "dividends": co.cfg["cash"].get("dividends", [])},
        "audit": audit_result, "files": files,
    }


def run(month=None, recalc=True, out_root=None, root=None, progress=None, slug=None, companies_root=None):
    """root: override of the company's actuals folder (tests); companies_root: override of companies/."""
    step = progress or (lambda s: None)
    co = company.load(slug, companies_root)
    month = month or co.cfg["company"]["last_closed_month"]
    step("actuals")
    from . import ledger
    if month not in ledger.available_months(co, root):
        from .variance import ReportError
        raise ReportError("no actuals for %s" % month)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    out = Path(out_root or OUTPUT) / co.slug / ("%s_%s" % (month, stamp))
    out.mkdir(parents=True, exist_ok=True)
    name = "%s_FPA_%s.xlsx" % (re.sub(r"\W+", "_", co.name.replace(" Ltd", "")).strip("_"), month)
    path = out / name
    step("workbook")
    lay = workbook.build(co, path, month=month, root=root)
    audit_result = None
    if recalc:
        step("recalc")
        from .recalc import recalc as excel_recalc
        excel_recalc(path)
        step("audit")
        audit_result = audit(co, path, lay, root)
    data = summary(co, lay, audit_result, {"workbook": name}, root)
    (out / "summary.json").write_text(json.dumps(data, indent=1, default=str), encoding="utf-8")
    return {"dir": str(out), "workbook": str(path), "summary": data}

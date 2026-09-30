"""Phase 6: management pack (Cover, Dashboard, Checks), monthly pipeline and audit."""
import json

import openpyxl
import pytest

from fpa import pipeline


def test_pipeline_without_excel_writes_pack_and_summary(tmp_path):
    res = pipeline.run("2026-09", recalc=False, out_root=tmp_path)
    wb = openpyxl.load_workbook(res["workbook"])
    assert wb.sheetnames[:3] == ["Cover", "Dashboard", "BvA"] and wb.sheetnames[-1] == "Checks"
    assert len(wb["Dashboard"]._charts) == 3
    s = json.loads((tmp_path / res["dir"].split("/")[-1] / "summary.json").read_text())
    assert s["month"] == "2026-09" and s["forecast_label"] == "6+6"
    assert len(s["monthly"]) == 12 and len(s["cash"]["closing"]) == 13
    assert s["full_year"]["downside"]["ebitda"] < s["full_year"]["base"]["ebitda"] < s["full_year"]["upside"]["ebitda"]
    assert s["commentary"]["ytd"][0]["label"] == "Headline"
    assert s["audit"] is None


def test_pipeline_rejects_a_month_without_actuals(tmp_path):
    with pytest.raises(ValueError, match="no actuals"):
        pipeline.run("2026-11", recalc=False, out_root=tmp_path)


@pytest.mark.excel
def test_pipeline_audit_is_clean(tmp_path):
    res = pipeline.run("2026-09", out_root=tmp_path)
    au = res["summary"]["audit"]
    assert au["values_checked"] > 400
    assert au["mismatches"] == [] and au["excel_errors"] == [] and au["model_checks"] == "OK"

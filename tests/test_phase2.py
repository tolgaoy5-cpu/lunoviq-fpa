"""Phase 2: driver-based budget, Python engine vs the formula-driven Excel sheets."""
import openpyxl
import pytest

from fpa import budget, company, workbook

CO = company.load()


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    path = tmp_path_factory.mktemp("wb") / "budget.xlsx"
    return path, workbook.build(CO, path)


def test_budget_totals_are_consistent():
    b = budget.build(CO)
    fy, t = budget.full_year(b)
    assert t["revenue"] == pytest.approx(sum(b["totals"][m]["revenue"] for m in b["months"]))
    assert t["ebitda"] == pytest.approx(t["revenue"] - t["opex"])
    assert 0.15 < t["ebitda"] / t["revenue"] < 0.25


def test_every_budget_figure_is_a_formula(built):
    path, lay = built
    ws = openpyxl.load_workbook(path)["Budget"]
    for code in CO.chart.by_code:
        row = lay["budget"]["L" + code]
        for i in range(12):
            v = ws["%s%d" % (workbook.mcol(i), row)].value
            assert isinstance(v, str) and v.startswith("="), (code, i, v)


def test_inputs_are_named_and_blue(built):
    path, _ = built
    wb = openpyxl.load_workbook(path)
    for sec, key, *_ in workbook.SCALARS:
        dn = wb.defined_names["a_" + key]
        sheet, ref = next(dn.destinations)
        cell = wb[sheet][ref.replace("$", "")]
        assert cell.font.color.rgb.endswith("0000FF")
        section = {"Portfolio": "portfolio", "Fees": "fees", "Staff": "staff", "Costs": "costs"}[sec]
        assert cell.value == CO.cfg[section][key]


@pytest.mark.excel
def test_excel_budget_matches_python(built):
    from fpa.recalc import recalc
    path, lay = built
    recalc(path)
    ws = openpyxl.load_workbook(path, data_only=True)["Budget"]
    py = budget.build(CO)
    for i, m in enumerate(lay["months"]):
        col = workbook.mcol(i)
        for code in CO.chart.by_code:
            assert ws["%s%d" % (col, lay["budget"]["L" + code])].value == pytest.approx(py["pnl"][m][code], abs=0.05)
        for k in ("revenue", "opex", "ebitda", "ebit", "net_income"):
            assert ws["%s%d" % (col, lay["budget"][k])].value == pytest.approx(py["totals"][m][k], abs=0.1)
    fy, t = budget.full_year(py)
    assert ws["%s%d" % (workbook.TOTAL, lay["budget"]["ebitda"])].value == pytest.approx(t["ebitda"], abs=0.5)

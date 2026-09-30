"""Any company: chart of accounts, importer (real-world export shapes), rule-based budget,
run-rate forecast, flexed-budget variance and the full pack for Brightwell Cleaning."""
import shutil

import openpyxl
import pytest

from fpa import accounts, budget, cash, company, forecast, importer, ledger, rules, variance, workbook
from fpa import demo_brightwell as demo

BW = "brightwell-cleaning"


@pytest.fixture()
def root(tmp_path):
    """A private copy of companies/ so tests never touch the real data."""
    shutil.copytree(company.COMPANIES, tmp_path / "companies")
    return tmp_path / "companies"


@pytest.fixture(scope="module")
def bw():
    return company.load(BW)


# ---------------------------------------------------------------- chart
@pytest.mark.parametrize("rows, msg", [
    ([("1", "A", "L", "revenue"), ("1", "B", "L", "opex")], "duplicate"),
    ([("1", "A", "L", "revenue"), ("2", "B", "L", "costs")], "unknown section"),
    ([("2", "B", "L", "opex")], "no revenue"),
])
def test_chart_validation(rows, msg):
    with pytest.raises(accounts.ChartError, match=msg):
        accounts.Chart([accounts.Account(*r) for r in rows])


def test_chart_totals_with_cost_of_sales_and_interest(bw):
    p = {c: 0.0 for c in bw.chart.by_code}
    p.update({"200": 1000, "310": 400, "477": 300, "416": 50, "437": 20, "500": 57.5})
    t = bw.chart.totals(p)
    assert (t["gross_profit"], t["ebitda"], t["ebit"], t["pbt"], t["net_income"]) == (600, 300, 250, 230, 172.5)
    assert bw.chart.has_cost_of_sales and bw.chart.has_interest
    assert "090" in bw.chart.all_by_code and "090" not in bw.chart.by_code        # balance sheet ignored


# ---------------------------------------------------------------- importer
def test_amounts_and_month_headers():
    assert importer.amount("£1,234.50") == 1234.5 and importer.amount("(12.00)") == -12
    assert importer.amount("") == 0 and importer.amount("12.5-") == -12.5
    assert [importer.month_of(h) for h in ("Apr-26", "April 2026", "2026-04", "30 Apr 2026", "Total")] == \
        ["2026-04"] * 4 + [None]
    with pytest.raises(importer.ImportError_):
        importer.amount("abc")


def test_trial_balance_with_titles_and_synonyms(bw):
    text = ("Trial Balance\nBrightwell\nAs at 30 April 2026\n\n"
            "Nominal Code;Nominal Name;Dr;Cr\n"
            "200;Sales - Commercial contracts;;1.000,00\n".replace("1.000,00", "1000.00")
            + "310;Direct wages - cleaners;400.00;\n090;Business bank account;600.00;\n")
    parsed = importer.parse(text)
    assert parsed["layout"] == "trial_balance" and parsed["header_line"] == 5
    pnl, info = importer.trial_balance_to_pnl(bw, parsed)
    assert pnl["200"] == 1000 and pnl["310"] == 400 and pnl["201"] == 0
    assert info["balanced"] is True and info["result"] == 600


def test_name_only_export_and_signed_balance(bw):
    text = "Account,Balance\nSales - Commercial contracts,-1000\nDirect wages - cleaners,400\n"
    pnl, _ = importer.trial_balance_to_pnl(bw, importer.parse(text))
    assert pnl["200"] == 1000 and pnl["310"] == 400


def test_unknown_accounts_are_named(bw):
    text = "Account Code,Account,Debit,Credit\n999,Mystery account,10,\n200,Sales,,10\n"
    with pytest.raises(importer.ImportError_, match="999 Mystery account"):
        importer.trial_balance_to_pnl(bw, importer.parse(text))


def test_ytd_needs_earlier_months(bw, root):
    co = company.load(BW, root)
    shutil.rmtree(co.actuals_dir / "2026-05")
    tb = (co.dir / "imports" / "trial_balance_2026-06.csv").read_text()
    with pytest.raises(importer.ImportError_, match="2026-05"):
        importer.import_month(co, "2026-06", tb)


def test_pnl_by_month_skips_headings_and_subtotals(bw):
    text = (bw.dir / "imports" / "pnl_by_month_FY2025-26.csv").read_text()
    parsed = importer.parse(text)
    assert parsed["layout"] == "pnl_by_month" and len(parsed["months"]) == 12
    assert not any(r["name"].startswith(("Total", "Gross", "Net")) for r in parsed["rows"])


def test_demo_exports_import_exactly(root):
    """Exports built like an accounting system's (titles, subtotals, YTD trial balances with the balance
    sheet) must import to exactly the amounts they were built from."""
    co = company.load(BW, root)
    shutil.rmtree(co.actuals_dir, ignore_errors=True)
    true = demo.build(co)
    for m, p in true.items():
        got, _ = ledger.read_month(co, m)
        assert got == pytest.approx(p, abs=0.005), m


# ---------------------------------------------------------------- rules budget
def test_rules_budget(bw):
    b = budget.build(bw)
    rs = rules.rules(bw)
    py = {m: ledger.read_month(bw, m)[0] for m in bw.fy_months(2025)}
    for i, m in enumerate(b["months"]):
        p, pm = b["pnl"][m], py[bw.fy_months(2025)[i]]
        rev = sum(p[c] for c in bw.chart.codes("revenue"))
        assert p["200"] == pytest.approx(pm["200"] * 1.06, abs=0.01)                  # growth keeps seasonality
        assert p["310"] == pytest.approx(0.47 * rev, abs=0.02)                         # share of revenue
        assert p["469"] == 1650.0                                                      # fixed
        assert p["416"] == pm["416"]                                                   # prior year
        assert p["500"] == pytest.approx(0.25 * bw.chart.totals(p)["pbt"], abs=0.02)   # tax on PBT
    assert rs["311"]["method"] == "pct_revenue" and rs["449"] == {"method": "growth", "rate": 0.04}


def test_rules_need_last_years_actuals(root):
    co = company.load(BW, root)
    shutil.rmtree(co.actuals_dir / "2025-12")
    with pytest.raises(rules.RulesError, match="2025-12"):
        budget.build(co)


def test_bad_rule_is_rejected(bw):
    co = company.Company({**bw.cfg, "budget": {"accounts": {"200": {"method": "pct_revenue", "rate": 0.1}}}},
                         bw.slug, bw.dir, bw.chart)
    with pytest.raises(rules.RulesError, match="revenue account"):
        rules.rules(co)


# ---------------------------------------------------------------- forecast, variance, cash
def test_run_rate_forecast_and_scenarios(bw):
    sc = forecast.scenarios(bw, "2026-09")
    base = sc["base"]
    act_rev = sum(base["totals"][m]["revenue"] for m in base["closed"])
    b = budget.build(bw)
    bud_rev = sum(b["totals"][m]["revenue"] for m in base["closed"])
    assert base["assumptions"]["revenue_factor"] == pytest.approx(act_rev / bud_rev)
    m = base["open"][0]
    assert base["pnl"][m]["200"] == pytest.approx(b["pnl"][m]["200"] * act_rev / bud_rev, abs=0.02)
    assert base["pnl"][m]["320"] == pytest.approx(b["pnl"][m]["320"] * base["assumptions"]["cost_of_sales_factor"]
                                                  + 350, abs=0.02)
    e = {n: r["fy_totals"]["ebitda"] for n, r in sc.items()}
    assert e["downside"] < e["base"] < e["upside"]


def test_flexed_budget_and_always_threshold(bw):
    r = variance.analyse(bw, "2026-09")
    for period in ("month", "ytd"):
        for line, b in r["bridges"][period].items():
            assert b["volume"] + b["rate"] + b["other"] == pytest.approx(b["var"])
    lab = r["bridges"]["ytd"]["Direct labour"]
    assert lab["rate"] < -5000 and lab["volume"] > 0          # share up, sales down
    assert r["lines"]["ytd"]["200"]["flag"] and abs(r["lines"]["ytd"]["200"]["pct"]) < 0.05   # size alone
    texts = " ".join(t for _, t, _ in r["commentary"]["ytd"])
    assert "cost share" in texts and "office contract ended" in texts and "gross margin" in texts


def test_cash_timing_for_a_rules_company(bw):
    cf = cash.build(bw, "2026-09")
    assert cf["receipt_lines"] == ["Commercial contracts", "Domestic cleaning", "End of tenancy", "Specialist cleaning"]
    assert "Office rent" not in cf["payment_lines"] and "Bank interest" in cf["payment_lines"]
    assert cf["table"][("Payments", "Dividend")] == pytest.approx([0.0] * 8 + [-30000.0] + [0.0] * 4)
    # invoiced contract income arrives 30 days later: September invoices land from week 1
    assert cf["table"][("Receipts", "Commercial contracts")][0] > 0


@pytest.mark.excel
def test_brightwell_pack_audits_clean(tmp_path):
    from fpa import pipeline
    res = pipeline.run("2026-09", out_root=tmp_path, slug=BW)
    au = res["summary"]["audit"]
    assert au["values_checked"] > 400 and au["mismatches"] == [] and au["excel_errors"] == []
    assert au["model_checks"] == "OK"
    wb = openpyxl.load_workbook(res["workbook"])
    assert "Prior year" in wb.sheetnames and "Drivers" not in wb.sheetnames
    ws = wb["Budget"]
    r = workbook.pnl_rows(company.load(BW).chart, 5)
    assert ws["D%d" % r["L200"]].value.startswith("='Prior year'!D")
    assert ws["D%d" % r["L310"]].value.startswith("=Assumptions!$E$")

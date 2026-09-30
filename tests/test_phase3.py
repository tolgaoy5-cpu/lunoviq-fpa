"""Phase 3: budget vs actual, driver bridges, flags, commentary, Actuals and BvA sheets."""
import openpyxl
import pytest

from fpa import accounts, company, variance, workbook

CO = company.load()


@pytest.fixture(scope="module")
def res():
    return variance.analyse(CO, "2026-09")


def test_variance_signs_and_totals(res):
    for period in ("month", "ytd"):
        for code, l in res["lines"][period].items():
            expected = l["actual"] - l["budget"] if accounts.BY_CODE[code].kind == "revenue" else l["budget"] - l["actual"]
            assert l["var"] == pytest.approx(expected)
        t = res["totals"][period]
        assert t["ebitda"]["var"] == pytest.approx(t["revenue"]["var"] + t["opex"]["var"])
    assert res["totals"]["ytd"]["revenue"]["actual"] == pytest.approx(
        sum(res["actuals"][m]["pnl"][c] for m in res["ytd_months"] for c in accounts.REVENUE))


def test_bridges_add_up_to_the_variance(res):
    for period in ("month", "ytd"):
        for name, b in res["bridges"][period].items():
            assert b["volume"] + b["rate"] + b["other"] == pytest.approx(b["var"])
    staff = res["bridges"]["ytd"]["Staff costs"]
    assert staff["volume"] > 0                 # the vacancy saved salary cost
    assert res["bridges"]["ytd"]["Management fees"]["rate"] > 0      # rents ahead of budget


def test_flags_and_commentary_tell_the_story(res):
    ytd = dict((label, text) for label, text, _ in res["commentary"]["ytd"])
    assert "Let-only fees" in ytd and "let-only lets vs" in ytd["Let-only fees"]
    assert "administration" in ytd["Bad debts"]
    assert "Agency cover" in ytd["Other overheads"]
    assert res["commentary"]["month"][0][0] == "Headline"
    assert res["lines"]["ytd"]["6600"]["flag"] and not res["lines"]["ytd"]["6300"]["flag"]


def test_reporting_month_must_be_closed():
    with pytest.raises(variance.ReportError, match="missing"):
        variance.analyse(CO, "2026-10")
    with pytest.raises(variance.ReportError, match="not in"):
        variance.analyse(CO, "2025-09")


@pytest.mark.excel
def test_excel_bva_matches_python(tmp_path):
    from fpa.recalc import recalc
    path = tmp_path / "pack.xlsx"
    lay = workbook.build(CO, path, month="2026-09")
    recalc(path)
    ws = openpyxl.load_workbook(path, data_only=True)["BvA"]
    res = lay["variance"]
    for code in accounts.BY_CODE:
        r = lay["bva"][code]
        for col, per, f in (("D", "month", "actual"), ("E", "month", "budget"), ("F", "month", "var"),
                            ("I", "ytd", "actual"), ("J", "ytd", "budget"), ("K", "ytd", "var")):
            assert ws["%s%d" % (col, r)].value == pytest.approx(res["lines"][per][code][f], abs=0.1)
        flag = ("M " if res["lines"]["month"][code]["flag"] else "") + ("YTD" if res["lines"]["ytd"][code]["flag"] else "")
        assert (ws["M%d" % r].value or "") == flag.strip()
    for k in ("revenue", "opex", "ebitda", "net_income"):
        assert ws["K%d" % lay["bva"][k]].value == pytest.approx(res["totals"]["ytd"][k]["var"], abs=0.5)

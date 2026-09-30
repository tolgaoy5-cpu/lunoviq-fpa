"""Phase 4: rolling forecast (6+6) and scenarios."""
import openpyxl
import pytest

from fpa import accounts, company, forecast, ledger, workbook

CO = company.load()


@pytest.fixture(scope="module")
def sc():
    return forecast.scenarios(CO, "2026-09")


def test_reforecast_combines_actuals_and_forecast(sc):
    fc = sc["base"]
    assert fc["label"] == "6+6" and len(fc["closed"]) == 6 and len(fc["open"]) == 6
    for m in fc["closed"]:
        assert fc["pnl"][m] == ledger.read_month(m)[0]
    assert fc["fy_totals"]["revenue"] == pytest.approx(sum(fc["totals"][m]["revenue"] for m in fc["months"]))
    first = fc["drivers"][0]
    k = ledger.read_month("2026-09")[1]
    assert first["pum_open"] == k["pum_close"]                       # starts from the actual position
    assert first["avg_rent"] == pytest.approx(k["avg_rent"] * 1.005)
    assert first["params"]["portals"] == 3680.0                      # latest estimate applied
    assert 0.6 <= fc["lets_factor"] < 1.0                              # YTD lets below budget carried forward


def test_scenarios_are_ordered(sc):
    e = {n: r["fy_totals"]["ebitda"] for n, r in sc.items()}
    assert e["downside"] < e["base"] < e["upside"]
    assert all(r["pnl"]["2026-04"] == sc["base"]["pnl"]["2026-04"] for r in sc.values())   # actuals unchanged
    assert sc["base"]["fy_totals"]["ebitda"] < sc["base"]["budget_fy_totals"]["ebitda"]


def test_forecast_errors():
    with pytest.raises(forecast.ForecastError):
        forecast.reforecast(CO, "2027-06")
    with pytest.raises(forecast.ForecastError, match="unknown scenario"):
        forecast.reforecast(CO, "2026-09", scenario="moonshot")


@pytest.mark.excel
def test_excel_forecast_matches_python(tmp_path):
    from fpa.recalc import recalc
    path = tmp_path / "pack.xlsx"
    lay = workbook.build(CO, path, month="2026-09")
    recalc(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, r = wb["Forecast"], lay["forecast"]["pnl"]
    fc = lay["scenario_results"]["base"]
    for code in accounts.BY_CODE:
        assert ws["%s%d" % (workbook.TOTAL, r["L" + code])].value == pytest.approx(fc["fy"][code], abs=0.1)
    for k in ("revenue", "ebitda", "net_income"):
        assert ws["%s%d" % (workbook.TOTAL, r[k])].value == pytest.approx(fc["fy_totals"][k], abs=0.5)
        assert ws["R%d" % r[k]].value == pytest.approx(fc["budget_fy_totals"][k], abs=0.5)
    assert not [c for s in wb for row in s.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("#")]
